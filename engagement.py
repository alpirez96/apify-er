"""Shortlist Excel'indeki influencer'lar için Instagram ve TikTok engagement rate raporu üretir.

Kullanım:
    python engagement.py "Samsung - Shortlist.xlsx" --sadece-linkler   # hesapları ve olası link hatalarını listeler
    APIFY_TOKEN=... python engagement.py "Samsung - Shortlist.xlsx"    # Apify'ı çalıştırıp raporu oluşturur
"""

import argparse
import os
import re
import sys
from collections import Counter
from datetime import date
from pathlib import Path

from apify_client import ApifyClient
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

POST_SAYISI = 5

# Shortlist düzeni (ilk sayfa): 2. satır başlık, A=Influencer, B=IG takipçi (profil linki bağlı), D=TT takipçi (link bağlı)
BASLIK_SATIRI = 2
ISIM_KOLON = "A"
PLATFORMLAR = {
    "Instagram": ("B", re.compile(r"instagram\.com/([A-Za-z0-9._]+)", re.I), "https://www.instagram.com/{}/"),
    "TikTok": ("D", re.compile(r"tiktok\.com/@([A-Za-z0-9._]+)", re.I), "https://www.tiktok.com/@{}"),
}

IG_ACTOR = "apify/instagram-profile-scraper"
TT_ACTOR = "clockworks/tiktok-scraper"

# Post etkileşimi = aşağıdaki alanların toplamı (Apify çıktısındaki alan adları).
# Instagram = Beğeni + Yorum, TikTok = Beğeni + Yorum + Paylaşım
IG_ALANLAR = {"Beğeni": "likesCount", "Yorum": "commentsCount"}
TT_ALANLAR = {"Beğeni": "diggCount", "Yorum": "commentCount", "Paylaşım": "shareCount"}

TR_ASCII = str.maketrans("ıİğĞüÜşŞöÖçÇ", "iIgGuUsSoOcC")


# --- Excel'den hesapları okuma ---

def sayi(deger):
    """'837K', '1.5M', 3600000 gibi takipçi değerlerini sayıya çevirir."""
    if isinstance(deger, (int, float)):
        return deger
    m = re.fullmatch(r"\s*([\d.]+)\s*([KM]?)\s*", str(deger or ""), re.I)
    return float(m.group(1)) * {"": 1, "K": 1e3, "M": 1e6}[m.group(2).upper()] if m else None


def hesaplari_bul(ws):
    """Her influencer için {no, isim, hesaplar: {platform: (kullanıcı adı, Excel'deki takipçi)}} döner."""
    satirlar = []
    for no in range(BASLIK_SATIRI + 1, ws.max_row + 1):
        hesaplar = {}
        for platform, (kolon, regex, _) in PLATFORMLAR.items():
            hucre = ws[f"{kolon}{no}"]
            eslesme = regex.search(f"{hucre.hyperlink.target if hucre.hyperlink else ''} {hucre.value or ''}")
            if eslesme:
                hesaplar[platform] = (eslesme.group(1).lower(), sayi(hucre.value))
        if hesaplar:
            satirlar.append({"no": no, "isim": str(ws[f"{ISIM_KOLON}{no}"].value or "").strip(), "hesaplar": hesaplar})
    return satirlar


# --- Link kontrolü ---

def sade(metin):
    return re.sub(r"[^a-z0-9]", "", str(metin or "").translate(TR_ASCII).lower())


def isim_uyuyor(isim, *hesap_adlari):
    """Influencer adındaki kelimelerden biri kullanıcı adında veya profil adında geçiyor mu?"""
    hedef = " ".join(sade(a) for a in hesap_adlari)
    return any(len(k) >= 3 and k in hedef for k in map(sade, isim.split()))


def link_uyarilari(satirlar, veri=None):
    """Yanlış girilmiş olabilecek linkler için {(satır no, platform): [uyarılar]} döner.
    veri verilmezse (Apify çalışmadan) sadece tekrar ve kullanıcı adı kontrolü yapılır."""
    tekrar = Counter((p, ad) for s in satirlar for p, (ad, _) in s["hesaplar"].items())
    uyarilar = {}
    for s in satirlar:
        for platform, (ad, excel_takipci) in s["hesaplar"].items():
            hesap = veri[platform].get(ad) if veri else None
            u = []
            if tekrar[(platform, ad)] > 1:
                u.append("aynı hesap başka bir influencer'da da var")
            if veri and not hesap:
                u.append("Apify'da bulunamadı")
            if not isim_uyuyor(s["isim"], ad, hesap and hesap["ad"]):
                u.append("hesap adı influencer adıyla uyuşmuyor")
            if hesap and excel_takipci and not 0.5 <= hesap["takipci"] / excel_takipci <= 2:
                u.append("takipçi Excel'dekinden çok farklı")
            uyarilar[(s["no"], platform)] = u
    return uyarilar


# --- Apify ---

def actor_calistir(client, actor, run_input):
    print(f"-> {actor} çalışıyor...")
    run = client.actor(actor).call(run_input=run_input)
    if run is None or run.status != "SUCCEEDED":
        sys.exit(f"{actor} başarısız oldu: {run and run.status}")
    return list(client.dataset(run.default_dataset_id).iterate_items())


def metrikler(post, alanlar):
    # Gizli beğeniler Apify'da -1 / None gelebilir, 0 sayılır.
    m = {ad: max(post.get(alan) or 0, 0) for ad, alan in alanlar.items()}
    return {**m, "etkilesim": sum(m.values())}


def instagram_verisi(client, kullanicilar):
    items = actor_calistir(client, IG_ACTOR, {"usernames": kullanicilar})
    sonuc = {}
    for p in items:
        if not p.get("username"):
            continue
        postlar = sorted(p.get("latestPosts") or [], key=lambda x: x.get("timestamp") or "", reverse=True)
        sonuc[p["username"].lower()] = {
            "ad": p.get("fullName"),
            "takipci": p.get("followersCount") or 0,
            "postlar": [
                {"tarih": x.get("timestamp"), "url": x.get("url"), **metrikler(x, IG_ALANLAR)}
                for x in postlar[:POST_SAYISI]
            ],
        }
    return sonuc


def tiktok_verisi(client, kullanicilar):
    items = actor_calistir(client, TT_ACTOR, {
        "profiles": kullanicilar,
        "profileScrapeSections": ["videos"],
        "profileSorting": "latest",
        # Sabitlenmiş (pinned) eski videolar da gelebildiği için fazladan çekip tarihe göre sıralıyoruz.
        "resultsPerPage": POST_SAYISI + 5,
        "shouldDownloadVideos": False,
        "shouldDownloadCovers": False,
        "shouldDownloadSubtitles": False,
        "shouldDownloadSlideshowImages": False,
    })
    sonuc = {}
    for v in items:
        yazar = v.get("authorMeta") or {}
        if not yazar.get("name"):
            continue
        hesap = sonuc.setdefault(yazar["name"].lower(), {"ad": yazar.get("nickName"), "takipci": yazar.get("fans") or 0, "postlar": []})
        hesap["postlar"].append({"tarih": v.get("createTimeISO"), "url": v.get("webVideoUrl"), **metrikler(v, TT_ALANLAR)})
    for hesap in sonuc.values():
        hesap["postlar"] = sorted(hesap["postlar"], key=lambda x: x["tarih"] or "", reverse=True)[:POST_SAYISI]
    return sonuc


def engagement_rate(hesap):
    """ER = Toplam Etkileşim / (Takipçi x Post Sayısı), yani post başına ortalama etkileşim / takipçi."""
    if not hesap or not hesap["postlar"] or not hesap["takipci"]:
        return None
    toplam = sum(p["etkilesim"] for p in hesap["postlar"])
    return toplam / (hesap["takipci"] * len(hesap["postlar"]))


# --- Rapor ---

KOYU = "1428A0"
PLATFORM_RENK = {"Instagram ER": "C13584", "TikTok ER": "111111"}
TAKIPCI_FORMAT = '[>=1000000]0.0,,"M";[>=1000]0.0,"K";0'
INCE = Side(style="thin", color="D9D9D9")
KENAR = Border(left=INCE, right=INCE, top=INCE, bottom=INCE)


def yazi(**kw):
    return Font(name="Calibri", **kw)


UYARI_DOLGU, UYARI_FONT = PatternFill("solid", fgColor="FDE2E1"), yazi(color="B91C1C")
LINK_FONT = yazi(color="1155CC", underline="single")


def baslik_satiri(ws, satir, basliklar, renk="E7EAF3", font_rengi="000000"):
    for i, b in enumerate(basliklar, 1):
        c = ws.cell(satir, i, b)
        c.font = yazi(bold=True, color=font_rengi)
        c.fill = PatternFill("solid", fgColor=renk)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = KENAR


def hucre(ws, satir, kolon, deger, format=None, link=None):
    c = ws.cell(satir, kolon, deger)
    c.border, c.font = KENAR, yazi()
    c.alignment = Alignment(vertical="center", wrap_text=isinstance(deger, str))
    if format:
        c.number_format = format
    if link:
        c.hyperlink, c.font = link, LINK_FONT
    return c


def sayfa_ayari(ws, baslik_satirlari):
    ws.freeze_panes = ws.cell(baslik_satirlari[-1] + 1, 2)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{baslik_satirlari[0]}:{baslik_satirlari[-1]}"


def ozet_sayfasi(ws, kaynak, satirlar, veri, uyarilar):
    son_kolon = 8
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=son_kolon)
    ws["A1"] = "Influencer Engagement Raporu"
    ws["A1"].font = yazi(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=KOYU)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=son_kolon)
    ws["A2"] = f"{kaynak}  ·  {date.today():%d.%m.%Y}  ·  {len(satirlar)} influencer  ·  Kaynak: Apify"
    ws["A2"].font = yazi(italic=True, color="666666")

    formul = {"Instagram": f"ER = ({' + '.join(IG_ALANLAR)}) / Takipçi", "TikTok": f"ER = ({' + '.join(TT_ALANLAR)}) / Takipçi"}
    aciklamalar = [
        ("Instagram ER", f"{formul['Instagram']}   (son {POST_SAYISI} içeriğin ortalaması)"),
        ("TikTok ER", f"{formul['TikTok']}   (son {POST_SAYISI} içeriğin ortalaması)"),
        ("Hesaplama", f"Son {POST_SAYISI} içeriğin etkileşimleri toplanır: ER = Toplam Etkileşim / (Takipçi × {POST_SAYISI}). "
                      f"Örnek: 100K takipçi, son {POST_SAYISI} içerikte toplam 25.000 etkileşim → "
                      f"25.000 / (100.000 × {POST_SAYISI}) = %5,00"),
        ("Takipçi", "Rapor tarihinde Apify'dan güncel olarak çekildi."),
        ("Kontrol", "Yanlış girilmiş olabilecek linkler: aynı hesap birden fazla influencer'da, hesap adı influencer "
                    "adıyla uyuşmuyor, takipçi Excel'dekinden çok farklı ya da hesap bulunamadı. Detay: 'Link Kontrol'."),
        ("Not", "Beğenileri gizlenmiş Instagram içeriklerinde beğeni 0 sayılır. İçerik bazında rakamlar: 'Post Detay'."),
    ]
    for i, (etiket, metin) in enumerate(aciklamalar, 4):
        ws.cell(i, 1, etiket).font = yazi(bold=True)
        ws.cell(i, 1).alignment = Alignment(vertical="top")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=son_kolon)
        ws.cell(i, 2, metin).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 2).font = yazi()
        if len(metin) > 110:  # birleştirilmiş hücrelerde satır yüksekliği otomatik ayarlanmıyor
            ws.row_dimensions[i].height = 30
        for k in range(1, son_kolon + 1):
            ws.cell(i, k).fill = PatternFill("solid", fgColor="F5F6FA")
        if etiket in PLATFORM_RENK:  # formül satırları: platform renginde etiket, kalın formül
            ws.cell(i, 1).fill = PatternFill("solid", fgColor=PLATFORM_RENK[etiket])
            ws.cell(i, 1).font = yazi(bold=True, color="FFFFFF")
            ws.cell(i, 2).font = yazi(bold=True, size=12)

    # Tablo: üstte platform adı ve ER formülü, altında kolon başlıkları
    ust = 4 + len(aciklamalar) + 1
    for bas, son, platform in ((2, 4, "Instagram"), (5, 7, "TikTok")):
        renk = PLATFORM_RENK[f"{platform} ER"]
        for satir, metin, font in ((ust, platform, yazi(bold=True, color="FFFFFF")),
                                   (ust + 1, formul[platform], yazi(italic=True, size=9, color="FFFFFF"))):
            ws.merge_cells(start_row=satir, start_column=bas, end_row=satir, end_column=son)
            c = ws.cell(satir, bas, metin)
            c.fill, c.font, c.alignment = PatternFill("solid", fgColor=renk), font, Alignment(horizontal="center")
    baslik_satiri(ws, ust + 2, ["Influencer", "Hesap", "Takipçi", "ER", "Hesap", "Takipçi", "ER", "Kontrol"])

    ilk_veri = satir = ust + 3
    for s in satirlar:
        hucre(ws, satir, 1, s["isim"]).font = yazi(bold=True)
        notlar = []
        for kolon, platform in ((2, "Instagram"), (5, "TikTok")):
            ad, _ = s["hesaplar"].get(platform, (None, None))
            hesap = veri[platform].get(ad)
            hucre(ws, satir, kolon, ad and f"@{ad}", link=ad and PLATFORMLAR[platform][2].format(ad))
            hucre(ws, satir, kolon + 1, hesap and hesap["takipci"], TAKIPCI_FORMAT)
            hucre(ws, satir, kolon + 2, engagement_rate(hesap), "0.00%")
            if uyarilar.get((s["no"], platform)):
                notlar.append(f"{platform}: {', '.join(uyarilar[(s['no'], platform)])}")
        kontrol = hucre(ws, satir, 8, "\n".join(notlar) or None)
        if notlar:
            kontrol.fill, kontrol.font = UYARI_DOLGU, UYARI_FONT
        satir += 1

    for kolon in ("D", "G"):
        ws.conditional_formatting.add(f"{kolon}{ilk_veri}:{kolon}{satir - 1}",
                                      ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="63BE7B"))
    for kolon, genislik in zip("ABCDEFGH", (26, 26, 11, 9, 26, 11, 9, 48)):
        ws.column_dimensions[kolon].width = genislik
    sayfa_ayari(ws, [ust, ust + 2])


def link_kontrol_sayfasi(ws, satirlar, veri, uyarilar):
    baslik_satiri(ws, 1, ["Influencer", "Platform", "Hesap", "Profil Adı (Apify)", "Excel'deki Takipçi", "Güncel Takipçi", "Uyarı"])
    kayitlar = [(s, p, ad, excel) for s in satirlar for p, (ad, excel) in s["hesaplar"].items()]
    kayitlar.sort(key=lambda k: not uyarilar[(k[0]["no"], k[1])])  # uyarılı olanlar üstte
    for satir, (s, platform, ad, excel_takipci) in enumerate(kayitlar, 2):
        hesap = veri[platform].get(ad)
        u = uyarilar[(s["no"], platform)]
        hucre(ws, satir, 1, s["isim"])
        hucre(ws, satir, 2, platform)
        hucre(ws, satir, 3, f"@{ad}", link=PLATFORMLAR[platform][2].format(ad))
        hucre(ws, satir, 4, hesap and hesap["ad"])
        hucre(ws, satir, 5, excel_takipci, TAKIPCI_FORMAT)
        hucre(ws, satir, 6, hesap and hesap["takipci"], TAKIPCI_FORMAT)
        c = hucre(ws, satir, 7, ", ".join(u) or "OK")
        if u:
            c.fill, c.font = UYARI_DOLGU, UYARI_FONT
    for kolon, genislik in zip("ABCDEFG", (26, 11, 26, 28, 13, 13, 50)):
        ws.column_dimensions[kolon].width = genislik
    sayfa_ayari(ws, [1])


def post_detay_sayfasi(ws, satirlar, veri):
    baslik_satiri(ws, 1, ["Influencer", "Platform", "Hesap", "Takipçi", "Tarih", "İçerik",
                          "Beğeni", "Yorum", "Paylaşım", "Etkileşim", "İçerik ER"])
    satir = 2
    for s in satirlar:
        for platform, (ad, _) in s["hesaplar"].items():
            hesap = veri[platform].get(ad)
            for p in (hesap or {}).get("postlar", []):
                degerler = [s["isim"], platform, f"@{ad}", hesap["takipci"], (p["tarih"] or "")[:10], None,
                            p.get("Beğeni"), p.get("Yorum"), p.get("Paylaşım"), p["etkilesim"],
                            p["etkilesim"] / hesap["takipci"] if hesap["takipci"] else None]
                for kolon, d in enumerate(degerler, 1):
                    hucre(ws, satir, kolon, d, {4: TAKIPCI_FORMAT, 11: "0.00%"}.get(kolon, "#,##0" if kolon > 6 else None))
                hucre(ws, satir, 6, "Aç", link=p["url"])
                satir += 1
    for kolon, genislik in zip("ABCDEFGHIJK", (26, 11, 26, 10, 12, 8, 11, 9, 10, 11, 10)):
        ws.column_dimensions[kolon].width = genislik
    sayfa_ayari(ws, [1])


def rapor_yaz(cikti, kaynak, satirlar, veri, uyarilar):
    wb = Workbook()
    ozet_sayfasi(wb.active, kaynak, satirlar, veri, uyarilar)
    wb.active.title = "Özet"
    link_kontrol_sayfasi(wb.create_sheet("Link Kontrol"), satirlar, veri, uyarilar)
    post_detay_sayfasi(wb.create_sheet("Post Detay"), satirlar, veri)
    wb.save(cikti)


def uyarilari_yazdir(satirlar, uyarilar):
    for s in satirlar:
        for platform, (ad, _) in s["hesaplar"].items():
            if uyarilar[(s["no"], platform)]:
                print(f"  UYARI #{s['no']} {s['isim']} | {platform} @{ad}: {', '.join(uyarilar[(s['no'], platform)])}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", type=Path)
    parser.add_argument("--sadece-linkler", action="store_true", help="Apify çalıştırmadan hesapları ve olası link hatalarını listele")
    args = parser.parse_args()

    satirlar = hesaplari_bul(load_workbook(args.excel).worksheets[0])
    print(f"{len(satirlar)} influencer bulundu:")
    for s in satirlar:
        ig, tt = (s["hesaplar"].get(p, ("-",))[0] for p in PLATFORMLAR)
        print(f"  #{s['no']:<3} {s['isim']:<28} instagram={ig:<26} tiktok={tt}")
    if args.sadece_linkler:
        uyarilari_yazdir(satirlar, link_uyarilari(satirlar))
        return

    token = os.environ.get("APIFY_TOKEN")
    if not token:
        sys.exit("APIFY_TOKEN ortam değişkeni tanımlı değil.")
    client = ApifyClient(token)

    kullanicilar = {p: sorted({s["hesaplar"][p][0] for s in satirlar if p in s["hesaplar"]}) for p in PLATFORMLAR}
    veri = {
        "Instagram": instagram_verisi(client, kullanicilar["Instagram"]) if kullanicilar["Instagram"] else {},
        "TikTok": tiktok_verisi(client, kullanicilar["TikTok"]) if kullanicilar["TikTok"] else {},
    }
    uyarilar = link_uyarilari(satirlar, veri)
    uyarilari_yazdir(satirlar, uyarilar)

    cikti = args.excel.with_name(f"{args.excel.stem} - ER Raporu.xlsx")
    rapor_yaz(cikti, args.excel.name, satirlar, veri, uyarilar)
    print(f"Kaydedildi: {cikti}")


if __name__ == "__main__":
    main()
